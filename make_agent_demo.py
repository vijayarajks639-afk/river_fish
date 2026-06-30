"""RF-19 — "How an AI Agent Team Built river_fish" — executive/stakeholder demo deck.

Tells the multi-agent DELIVERY story (not the product): the operating model, how work was
parallelized under governance, what the agents themselves said, and the repeatable method.
Re-run:  python make_agent_demo.py
"""
from __future__ import annotations

import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TODAY = datetime.date.today().strftime("%d %b %Y")
NAVY = RGBColor(0x00, 0x2B, 0x5B); GOLD = RGBColor(0xE6, 0xB8, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); LGRAY = RGBColor(0xF2, 0xF2, 0xF2)
DK = RGBColor(0x33, 0x33, 0x33); GREEN = RGBColor(0x0A, 0x7A, 0x33)
RED = RGBColor(0xB0, 0x00, 0x00); BLUE = RGBColor(0x1F, 0x5C, 0xA8)

prs = Presentation(); prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(s, l, t, w, h, c):
    sh = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.line.fill.background(); return sh


def txt(s, text, l, t, w, h, size=18, bold=False, color=DK, align=PP_ALIGN.LEFT, italic=False):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color; return tb


def bullets(s, items, l, t, w, h, size=13, color=DK, char="▸", gap=4):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.add_paragraph() if i else tf.paragraphs[0]; p.space_after = Pt(gap)
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(); r.text = f"{char}  {lead}"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = NAVY
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = color
        else:
            r = p.add_run(); r.text = f"{char}  {it}"; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def header(s, title, sub=None):
    rect(s, 0, 0, 13.33, 1.1, NAVY); txt(s, title, 0.3, 0.1, 12.7, 0.62, size=27, bold=True, color=WHITE)
    if sub: txt(s, sub, 0.3, 0.72, 12.7, 0.34, size=12.5, color=GOLD)
    rect(s, 0, 1.1, 13.33, 0.04, GOLD)


def label(s, text, l, t, w):
    rect(s, l, t, w, 0.3, NAVY); txt(s, text.upper(), l + 0.1, t + 0.03, w - 0.2, 0.24, size=11, bold=True, color=WHITE)


# 1 — TITLE
s = prs.slides.add_slide(BLANK); rect(s, 0, 0, 13.33, 7.5, NAVY); rect(s, 0, 3.05, 13.33, 0.06, GOLD)
txt(s, "How an AI Agent Team Built river_fish", 0.8, 1.7, 11.7, 1.0, size=38, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, "A multi-agent software-delivery story — speed with governance intact", 0.8, 2.95, 11.7, 0.5, size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
txt(s, "Column-Level Data Lineage Prototype  ·  Banking / on-prem  ·  " + TODAY, 0.8, 3.4, 11.7, 0.4, size=13, color=LGRAY, align=PP_ALIGN.CENTER)
txt(s, "Vijayaraj Shanmugam — Product Owner & Orchestration   |   Claude — Orchestrator + Specialist Agents",
    0.8, 6.7, 11.7, 0.4, size=11, color=RGBColor(0x9F, 0xB3, 0xC8), align=PP_ALIGN.CENTER)

# 2 — THE CHALLENGE
s = prs.slides.add_slide(BLANK); header(s, "The Challenge", "Build an audit-grade prototype fast — without losing control")
rect(s, 0.35, 1.35, 12.63, 1.5, LGRAY)
txt(s, "Build a credible, audit-grade column-level data-lineage prototype for a multi-source banking "
       "pipeline — with a natural-language interface and a live demo — in a single working session, "
       "while keeping every governance control intact.", 0.6, 1.55, 12.1, 1.1, size=14)
label(s, "The tension every leader feels with AI-accelerated delivery", 0.35, 3.1, 12.63)
bullets(s, [
    ("Speed vs. assurance — ", "move fast, but the lineage must be auditable (BCBS 239), not a black box."),
    ("Parallelism vs. collisions — ", "many workstreams at once, but no stepping on each other's code."),
    ("AI reach vs. data governance — ", "use LLMs, but customer data cannot leave / cannot enter a prompt."),
    ("Throughput vs. quality — ", "compress time, but keep a real QA gate and honest reporting."),
], 0.6, 3.5, 12.1, 3.0, size=14, gap=8)

# 3 — THE TEAM / OPERATING MODEL
s = prs.slides.add_slide(BLANK); header(s, "The Team & Operating Model", "One human owner · one AI orchestrator · specialist AI agents · agile cadence")
rect(s, 0.35, 1.35, 12.63, 0.55, BLUE)
txt(s, "Product Owner (human)  →  Orchestrator (lead AI)  →  Specialist Agents (parallel)",
    0.5, 1.42, 12.3, 0.4, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rows = [
    ("Product Owner — human", "Vision, scope, sign-off, governance, credentials, push authority", GOLD),
    ("Orchestrator — lead AI", "Architecture, integration, M4 NL interface, QA gate, ALL git commits", BLUE),
    ("Data Generator agent", "Synthetic 11-system bank + DuckDB + extended SQL", GREEN),
    ("Research / Docs agent", "Market & future-of-data + learning/interview docs (cited)", GREEN),
    ("Reg-RAG UI agent", "Streamlit UI + deploy prep for the sister project", GREEN),
    ("QA + Deploy-prep agents", "Independent red-team review · HF deploy packaging", GREEN),
]
for i, (role, owns, c) in enumerate(rows):
    y = 2.1 + i * 0.82; rect(s, 0.35, y, 12.63, 0.72, LGRAY if i % 2 == 0 else WHITE); rect(s, 0.35, y, 0.1, 0.72, c)
    txt(s, role, 0.6, y + 0.13, 4.0, 0.5, size=13.5, bold=True, color=NAVY)
    txt(s, owns, 4.8, y + 0.13, 8.0, 0.5, size=12.5, color=DK)
txt(s, "Cadence: Jira story-first · live worklogs · interim QA checkpoint · sprint retro (agents included)",
    0.35, 7.0, 12.63, 0.35, size=11.5, italic=True, color=NAVY)

# 4 — HOW WORK WAS PARALLELIZED
s = prs.slides.add_slide(BLANK); header(s, "How Work Was Parallelized — Without Collisions", "Partition by file surface · only the orchestrator commits the shared repo")
label(s, "File-surface ownership (the anti-collision rule)", 0.35, 1.35, 6.1)
bullets(s, [
    "Data Generator → generate_data.py, sql_extended/, schema.py",
    "Docs → MARKET_AND_FUTURE.md, LEARNING_LINEAGE.md (new files)",
    "Reg-RAG UI → reg_rag_assistant/* (separate directory)",
    "QA → CHECKPOINT_REVIEW_M3M5.md (review only, no code)",
    "Deploy-prep → requirements.txt, DEPLOY_HF.md, README front-matter",
    "Orchestrator → lineage/nl/integration + EVERY commit",
], 0.45, 1.75, 6.1, 3.6, size=12.5, gap=6)
label(s, "Why it works", 6.65, 1.35, 6.33)
bullets(s, [
    ("No two writers per file — ", "true parallelism, zero merge conflicts."),
    ("Single committer — ", "the orchestrator serializes git; agents just produce files."),
    ("Cold-start briefs — ", "each agent reads-first (mandatory STEP 0) so it grounds in real code."),
    ("Continuous briefing log — ", "a living charter keeps agents informed as the design evolves."),
    ("Background execution — ", "agents run async while the orchestrator does critical-path work."),
], 6.75, 1.75, 6.2, 3.6, size=12.5, gap=8)
rect(s, 0.35, 5.6, 12.63, 1.4, RGBColor(0xE8, 0xF0, 0xFA))
txt(s, "Time compression", 0.5, 5.7, 3.0, 0.35, size=13, bold=True, color=BLUE)
txt(s, "Four workstreams advanced concurrently instead of sequentially; the orchestrator integrated while "
       "specialists built. The binding constraint was clarity of brief — not headcount, not tooling.",
    0.5, 6.05, 12.3, 0.9, size=13, color=DK)

# 5 — GOVERNANCE UNDER PARALLELISM
s = prs.slides.add_slide(BLANK); header(s, "Governance Under Parallelism", "The hard part of AI delivery — controls that held across every agent")
items = [
    ("Synthetic vs. real data boundary — ", "customer data 100% synthetic (DPDP-safe); only PUBLIC market data is real. Held across all agents."),
    ("Audit-grade backbone — ", "lineage is deterministic (parser → DAG); the LLM only explains and is labelled 'suggestion — needs sign-off'."),
    ("On-prem proof, mechanical — ", "the eval harness asserts the engine reads SQL text only, imports no data engine, and no rows enter a prompt."),
    ("Human-in-the-loop control points — ", "sign-off, repository pushes, and ALL credential handling stayed with the human."),
    ("Honest QA — ", "an interim red-team gate; failures reported, never hidden; the agents' own retro caught a security risk before harm."),
]
for i, (lead, rest) in enumerate(items):
    y = 1.45 + i * 1.05; rect(s, 0.35, y, 12.63, 0.92, LGRAY if i % 2 == 0 else WHITE); rect(s, 0.35, y, 0.1, 0.92, GREEN)
    bullets(s, [(lead, rest)], 0.6, y + 0.16, 12.2, 0.7, size=13.5)

# 6 — WHAT THE AGENTS SAID
s = prs.slides.add_slide(BLANK); header(s, "What the Agents Said — Retro Voices", "We ran the retro WITH the agents (one team, not human-vs-AI)")
voices = [
    ("Reg-RAG UI agent", "“Clear briefing meant zero rework.”", RED,
     "Security catch: a PUBLIC demo Space with an API key lets anyone spend it — keep public demos $0/key-less."),
    ("Research / Docs agent", "“Read-first briefing gave the exact repo vocabulary.”", GOLD,
     "Caution: vendor/Gartner claims age fast — say ‘analysts suggest’, not ‘Gartner says’."),
    ("Data Generator agent", "“The clean pipeline contract let me extend confidently.”", GREEN,
     "Flagged its own duplicate-key shortcut → orchestrator merged it before commit."),
]
for i, (who, quote, c, risk) in enumerate(voices):
    y = 1.45 + i * 1.78; rect(s, 0.35, y, 12.63, 1.62, LGRAY); rect(s, 0.35, y, 0.12, 1.62, c)
    txt(s, who, 0.6, y + 0.12, 4.2, 0.4, size=14, bold=True, color=NAVY)
    txt(s, quote, 0.6, y + 0.55, 5.6, 0.9, size=13, italic=True, color=DK)
    txt(s, "⚑ " + risk, 6.4, y + 0.2, 6.4, 1.2, size=12.5, color=DK)

# 7 — OUTCOMES & METRICS
s = prs.slides.add_slide(BLANK); header(s, "Outcomes & Metrics", "Shipped, verified, governed — in one session")
cards = [
    ("7", "milestones (M1–M7)", GREEN), ("18/18", "regression tests", GREEN),
    ("20/20", "eval (correctness+governance)", GREEN), ("11", "source systems modelled", BLUE),
    ("18,742", "synthetic rows (zero PII)", BLUE), ("~17", "commits, all pushed", BLUE),
    ("5", "AI agents on the team", GOLD), ("~9.5h", "AI-actual, 1 session", GOLD),
]
for i, (big, small, c) in enumerate(cards):
    col = i % 4; row = i // 4; x = 0.45 + col * 3.15; y = 1.5 + row * 1.9
    rect(s, x, y, 2.95, 1.7, LGRAY); rect(s, x, y, 2.95, 0.12, c)
    txt(s, big, x, y + 0.25, 2.95, 0.7, size=30, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    txt(s, small, x + 0.1, y + 1.0, 2.75, 0.6, size=12, color=DK, align=PP_ALIGN.CENTER)
txt(s, "Quality gates green throughout · governance proven mechanically · honest QA with failures reported",
    0.45, 5.5, 12.5, 0.4, size=12.5, italic=True, color=NAVY, align=PP_ALIGN.CENTER)
rect(s, 0.45, 6.0, 12.43, 1.0, RGBColor(0xE8, 0xF4, 0xE8))
txt(s, "The leadership takeaway: a repeatable, governed method to accelerate delivery with an AI agent "
       "team — durable briefs, file-surface partitioning, retros that include the agents, and the human "
       "owning sign-off. One team, not human-vs-AI.", 0.6, 6.12, 12.1, 0.85, size=13, color=DK)

# 8 — THE REPEATABLE METHOD
s = prs.slides.add_slide(BLANK); header(s, "The Repeatable Method", "What to reuse on the next prototype")
bullets(s, [
    ("Brief agents cold, read-first — ", "a mandatory STEP 0 to read the real code grounds output and kills rework."),
    ("Partition by file surface — ", "each agent owns distinct files; one committer serializes the repo."),
    ("Keep a living briefing log — ", "‘continuously informed’ charters so agents absorb design changes."),
    ("Run async, integrate centrally — ", "specialists in the background; the orchestrator on the critical path."),
    ("Gate with honest QA — ", "an independent red-team; report failures; fix before downstream impact."),
    ("Retro WITH the agents — ", "their candid feedback surfaced a real security risk and a latent bug."),
    ("Human owns the irreversibles — ", "sign-off, pushes, credentials, and the data-governance boundary."),
], 0.5, 1.45, 12.4, 4.6, size=14, gap=10)
rect(s, 0.35, 6.55, 12.63, 0.6, NAVY)
txt(s, "river_fish · github.com/vijayarajks639-afk/river_fish · synthetic data · real public market data · one team",
    0.5, 6.66, 12.3, 0.4, size=11.5, color=WHITE, align=PP_ALIGN.CENTER)

out = r"c:\Users\vijay\claude\Projects\river_fish\Agent_Team_Demo_river_fish.pptx"
prs.save(out); print("Saved:", out)
