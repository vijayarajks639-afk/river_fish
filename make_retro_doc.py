"""Generate the executive-grade 'AI Agent Team — Sprint 1 Retrospective' Word document.

For Technology Executive / stakeholder presentations. Records each specialist agent's
candid retro feedback plus a cross-cutting synthesis and an executive 'why it matters'.

Re-run any time the retro content changes:  python make_retro_doc.py
"""
from __future__ import annotations

import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

TODAY = datetime.date.today().strftime("%d %b %Y")

NAVY = RGBColor(0x00, 0x2B, 0x5B)
GOLD = RGBColor(0xB8, 0x86, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
RED  = RGBColor(0xB0, 0x00, 0x00)

# ── Agent retro voices (verbatim-faithful, lightly tidied) ────────────────────

AGENTS = [
    {
        "name": "Data Generator agent",
        "scope": "Synthetic multi-source banking data + DuckDB warehouse + sql_extended/ transforms",
        "output": "18,742 rows · 15 tables · 11 source systems · KNOWN_SCHEMA 9→34 tables · core tests 18/18 green",
        "voice": [
            ("What went well",
             "The existing pipeline contract was unusually clean to build against — explicit column lists in every "
             "sql/*.sql, a schema.py catalog already in place, and a test suite asserting exact graph shape "
             "(54/49/18). That allowed confident extension: all new transforms went in sql_extended/ so the core "
             "lineage glob never saw them, and the core tests stayed 18/18 green on the first run. Seeded "
             "determinism (SEED=42) and the FX join-key design (synthetic currency → real market spot rate → "
             "derived INR exposure) both landed first try."),
            ("Friction",
             "The main tension: the brief said 'register every new table in KNOWN_SCHEMA', but three source "
             "tables already existed there with their original columns. Editing the originals risked touching a "
             "contract the tests depend on; appending fresh entries was safer — the brief didn't specify which. "
             "Secondary: gap_analysis.py only globs sql/*.sql, so the ~35% intentionally-unconsumed columns in "
             "the extended systems don't surface in its report, though the brief implied they would."),
            ("What would have helped",
             "A one-line note on how KNOWN_SCHEMA should handle already-existing tables (extend-in-place vs. "
             "append) would have removed the judgment call. Also a stated decision on whether gap_analysis.py "
             "should be pointed at sql_extended/ too — right now the richest gap findings are invisible to the "
             "tool meant to report them."),
            ("Risk flagged",
             "Duplicate-key entries were appended for the three active source tables into KNOWN_SCHEMA rather "
             "than merged. Python keeps the last definition, so the extended supersets win — tests pass and "
             "behaviour is correct today — but a duplicate literal key is a latent foot-gun: it reads as an "
             "accident, and editing the first (now-dead) entry would silently have no effect. Recommend the "
             "orchestrator merge each pair before committing. Two smaller follow-ups: widen gap_analysis.py's "
             "glob if extended-system gaps should show; and the FX rate depends on a live fetch, so on an "
             "offline demo machine it falls back to 1.0 and INR exposure silently equals local exposure — worth "
             "a visible 'STALE' caption if shown on stage. Governance and referential integrity are clean: zero "
             "PII, synthetic markers throughout, every counterparty_id resolves to the master spine."),
        ],
    },
    {
        "name": "Research / Docs agent",
        "scope": "Market & future-of-data synthesis + learning/interview notes (MARKET_AND_FUTURE.md, LEARNING_LINEAGE.md)",
        "output": "2 cited docs · ~24 web-verified sources · directional claims explicitly labelled",
        "voice": [
            ("What went well",
             "The read-first briefing was the right call — reading the eight source files cold gave the exact "
             "internal vocabulary (the find_node bug, the deterministic/LLM split, the market-data governance "
             "boundary), so the docs cite the repo accurately rather than generically. Web verification was "
             "high-yield: nearly every market claim resolved to a primary or near-primary source, including the "
             "Spider 2.0 ICLR numbers and Snowflake's 'nothing leaves the governance boundary' language verbatim."),
            ("Friction",
             "Primary analyst sources are paywalled — Gartner Magic Quadrant placements and the consolidation "
             "forecast only surfaced via vendor blogs (Atlan, Ataccama) that have an incentive to frame the "
             "quadrant favourably. Market-size figures also conflicted across segment definitions "
             "(column-level lineage vs lineage-automation vs lineage-as-a-service are three different numbers); "
             "each was labelled by segment rather than blended."),
            ("What would have helped",
             "The brief was unusually clear — no real gap. The one accelerator would have been a stated "
             "'as-of' date convention up front, since 2025-dated sources were being reconciled against a "
             "June-2026 'today'; claims were dated inline instead."),
            ("Risk flagged",
             "Vendor and regulatory claims age fast — Gartner positions, Alation/Atlan product names, and MCP's "
             "governance-foundation status will likely shift within 6–12 months; re-verify before any high-stakes "
             "interview. The two 'directional' items rest on vendor-blog secondary sources — state them as "
             "'analysts suggest', not 'Gartner says', unless the primary document is pulled first."),
        ],
    },
    {
        "name": "Reg-RAG UI agent",
        "scope": "Streamlit UI + HuggingFace deploy prep for the sister Regulatory-RAG project",
        "output": "app.py (3 tabs) · HF front-matter · DEPLOY_HF.md · rag.py imported, not modified",
        "voice": [
            ("What went well",
             "Clean briefing meant zero rework. 'Read rag.py/ask.py first and import-don't-modify' kept the UI a "
             "thin layer over the existing pipeline; the venv and corpus were both present, so the task was "
             "unblocked end to end. The Streamlit 1.41 / use_container_width constraint was called out "
             "explicitly, which removed the one easy way to ship a broken Space."),
            ("Friction",
             "Minimal. A small ambiguity on the 'Corpus' tab: the index doesn't expose a per-document chunk map, "
             "so chunk_text() was re-run in the UI to list chunks — deterministic (same logic), but it means "
             "chunking runs twice. No environment or access friction."),
            ("What would have helped",
             "A one-line note on whether to rewrite requirements.txt vs. make a minimal edit would have saved a "
             "judgment call (minimal was chosen). For larger sprints, telling specialists which files are theirs "
             "vs. shared up front would prevent edit collisions — not an issue here as the UI files were net-new."),
            ("Risk flagged",
             "Secret handling on a public Space. The $0 fallback is safe, but if ANTHROPIC_API_KEY is added as a "
             "Space secret on a PUBLIC Space, anyone can run queries and spend the key — no auth, no rate limit. "
             "Keep the public demo key-less ($0 mode), or gate LLM mode behind a private Space / usage cap. "
             "Secondary: confirm the corpus is copied into the Space repo on deploy."),
        ],
    },
]

CROSS_CUTTING = {
    "went_well": [
        "Read-first briefs (a mandatory STEP 0) produced near-zero rework — every agent grounded itself in the "
        "actual code before acting.",
        "Explicit constraints up front prevented broken output — e.g. the Streamlit 1.41 compatibility rule and "
        "the synthetic-vs-real data governance boundary were stated, not assumed.",
        "File-surface partitioning (each agent owns distinct files; only the orchestrator commits the shared "
        "repo) meant true parallelism with no edit collisions.",
        "The environment was ready (venvs, corpus, live data) so agents were unblocked end to end.",
    ],
    "friction": [
        "Small unstated judgment calls (rewrite vs. minimal edit; an 'as-of' date convention; schema "
        "duplicate-key append vs. merge) — each minor, but each cost a decision.",
        "Primary analyst sources (Gartner) are paywalled; some market claims rest on vendor-blog secondaries.",
        "A couple of deterministic-but-duplicated computations (chunking re-run for display) — harmless, worth knowing.",
    ],
    "risks": [
        "SECURITY: a public demo Space with an API key set lets anyone spend the key — keep public demos $0 / "
        "key-less, or gate paid mode behind a private Space with a cap.",
        "Market/regulatory claims age fast — re-verify vendor and Gartner positions before high-stakes use.",
        "Schema hygiene: duplicate-key entries appended to KNOWN_SCHEMA should be merged into single entries "
        "(harmless today — last value wins — but a reviewer will flag it).",
    ],
    "actions": [
        "Keep public portfolio demos in $0 mode; never put a live key on a public Space.",
        "Add a house style for agent briefs: 'as-of' date convention + explicit shared-vs-owned file list.",
        "Merge the duplicate KNOWN_SCHEMA keys; re-run tests (expected to stay 18/18).",
        "Re-verify fast-aging market claims before any executive or interview use.",
    ],
}


# ── docx helpers ──────────────────────────────────────────────────────────────

def _set(run, size=11, bold=False, color=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def heading(doc, text, size=15, color=NAVY, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    _set(r, size=size, bold=True, color=color)
    return p


def body(doc, text, size=10.5, color=None, bold=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    _set(r, size=size, bold=bold, color=color)
    return p


def bullet(doc, text, size=10.5, color=None, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_lead:
        r = p.add_run(bold_lead)
        _set(r, size=size, bold=True, color=color or NAVY)
    r = p.add_run(text)
    _set(r, size=size, color=color)
    return p


def build():
    doc = Document()
    # base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # ── Title block ──
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("AI Agent Team — Sprint 1 Retrospective")
    _set(r, size=22, bold=True, color=NAVY)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("river_fish — Column-Level Data Lineage Prototype")
    _set(r, size=13, bold=True, color=GOLD)

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(f"Prepared for Technology Executive / Stakeholder Review   ·   {TODAY}")
    _set(r, size=10, color=GREY)

    m2 = doc.add_paragraph()
    m2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m2.add_run("ONE team — human product owner + AI orchestrator + specialist AI agents, on agile cadence")
    _set(r, size=10, italic=True, color=GREY)

    doc.add_paragraph()

    # ── 1. Executive summary ──
    heading(doc, "1.  Executive Summary")
    body(doc,
         "The river_fish prototype — a deterministic, audit-grade column-level data-lineage engine for a "
         "multi-source banking pipeline, with a governed LLM explanation layer and a natural-language "
         "interface — was built using a small team of specialised AI agents working in parallel under a human "
         "product owner and an AI orchestrator. The team ran on an agile cadence: stories, an interim QA gate, "
         "and this retrospective.")
    body(doc,
         "This document records each agent's candid retro feedback — what worked, what created friction, and "
         "the risks they raised — alongside a cross-cutting synthesis and an executive view of why the method "
         "matters. It is intended for reuse in stakeholder and leadership presentations.")
    body(doc,
         "Headline outcome: P1 hardening, the LLM explanation layer (M3), the natural-language interface (M4), "
         "two cited market/learning documents, live public market-data ingestion, and an 18,742-row synthetic "
         "multi-source dataset — delivered in a single working session with the parallel agent team, with the "
         "core regression suite green (18/18) throughout.",
         bold=False)

    # ── 2. The agent team ──
    heading(doc, "2.  The Agent Team")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for c, txt in zip(hdr, ("Role", "Owns", "Surface (no collisions)")):
        rr = c.paragraphs[0].add_run(txt); _set(rr, size=10, bold=True, color=NAVY)
    rows = [
        ("Product owner (human)", "Vision, scope, sign-off, governance, credentials", "Decisions — not files"),
        ("Orchestrator (lead AI)", "Architecture, integration, M4 NL interface, git, QA", "lineage.py, nl.py, schema validation, all commits"),
        ("Data Generator agent", "Synthetic multi-source data + DuckDB", "generate_data.py, sql_extended/, schema.py"),
        ("Research / Docs agent", "Market & learning documentation", "MARKET_AND_FUTURE.md, LEARNING_LINEAGE.md"),
        ("Reg-RAG UI agent", "Streamlit UI + deploy (sister project)", "reg_rag_assistant/* (separate dir)"),
    ]
    for role, owns, surface in rows:
        cells = table.add_row().cells
        for cell, txt, bold in ((cells[0], role, True), (cells[1], owns, False), (cells[2], surface, False)):
            rr = cell.paragraphs[0].add_run(txt); _set(rr, size=9.5, bold=bold)
    body(doc, "")
    bullet(doc, "Work was partitioned by file surface so agents ran truly in parallel; only the orchestrator "
                "commits the shared repository, eliminating git collisions.", size=9.5, color=GREY)

    # ── 3. Sprint outcomes ──
    heading(doc, "3.  Sprint 1 Outcomes")
    for txt in [
        "Stories closed: research, M1/M2 lineage engine, QA checkpoint, P1 hardening, M3 explanation, M4 NL interface.",
        "Quality: core regression suite 18/18 green throughout; one HIGH defect caught by QA before the NL layer depended on it.",
        "Live data: 9/9 public market symbols ingested fresh (FX, UST yields, VIX, equity indices).",
        "Synthetic data: 18,742 rows across 15 tables / 11 source systems; one lineage path runs real market feed → synthetic exposure → derived risk number.",
        "Governance: synthetic-vs-real data boundary preserved across every agent; lineage engine reads SQL text only; no data rows enter any LLM prompt.",
    ]:
        bullet(doc, txt)

    # ── 4. Agent retro voices ──
    heading(doc, "4.  Agent Retrospective Voices")
    body(doc, "Each specialist agent was asked the same four questions. Responses are recorded faithfully.",
         color=GREY, size=9.5)
    for a in AGENTS:
        heading(doc, a["name"], size=12.5, color=GOLD, space_before=8, space_after=2)
        body(doc, f"Scope: {a['scope']}", size=9.5, color=GREY)
        body(doc, f"Output: {a['output']}", size=9.5, color=GREY, space_after=4)
        for label, text in a["voice"]:
            bullet(doc, text, bold_lead=f"{label}: ")

    # ── 5. Cross-cutting ──
    heading(doc, "5.  Cross-Cutting Themes")
    heading(doc, "What went well", size=11.5, color=NAVY, space_before=6, space_after=2)
    for t in CROSS_CUTTING["went_well"]:
        bullet(doc, t)
    heading(doc, "Friction", size=11.5, color=NAVY, space_before=6, space_after=2)
    for t in CROSS_CUTTING["friction"]:
        bullet(doc, t)
    heading(doc, "Risks to note", size=11.5, color=RED, space_before=6, space_after=2)
    for t in CROSS_CUTTING["risks"]:
        bullet(doc, t, color=RED)
    heading(doc, "Actions", size=11.5, color=NAVY, space_before=6, space_after=2)
    for t in CROSS_CUTTING["actions"]:
        bullet(doc, t)

    # ── 6. Executive lens ──
    heading(doc, "6.  Why This Matters — Executive Lens")
    for lead, txt in [
        ("Time compression. ",
         "Four workstreams advanced concurrently instead of sequentially; the orchestrator integrated while "
         "specialists built. The constraint was clarity of brief, not headcount."),
        ("Governance under parallelism. ",
         "The hardest part of AI-accelerated delivery is keeping controls intact while moving fast. Here the "
         "synthetic-vs-real data boundary, the audit-grade deterministic backbone, and the 'LLM output is a "
         "suggestion, not a sign-off' rule held across every agent."),
        ("Human-in-the-loop control points. ",
         "Sign-off, repository pushes, and all credential handling stayed with the human. Agents proposed and "
         "produced; the human disposed. The retro itself surfaced a security risk (public key exposure) before "
         "any harm."),
        ("A repeatable method. ",
         "Durable charter documents, a continuously-updated briefing log ('keep agents informed'), file-surface "
         "partitioning, and retros that include the agents themselves — this is a method that scales to the "
         "next prototype, not a one-off."),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
        r = p.add_run(lead); _set(r, size=10.5, bold=True, color=NAVY)
        r = p.add_run(txt); _set(r, size=10.5)

    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run("Synthetic data only · Public market data is real and public · One team, not human-vs-AI")
    _set(r, size=9, italic=True, color=GREY)

    out = r"c:\Users\vijay\claude\Projects\river_fish\Multi_Agent_Retro_Sprint1.docx"
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    build()
